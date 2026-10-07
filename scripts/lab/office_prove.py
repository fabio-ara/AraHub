"""MAT06 local fixtures, native DOCX rendering and independent structure checks.

Run with the bundled Python and --config <private JSON>. Configuration supplies
output_dir, optional originals [{path, sha256}], and pinned tools
{docx_renderer, soffice, pdftoppm, pdfinfo}, each {path, sha256}. All paths are
absolute; output_dir must be inside this checkout's .private directory.
No source document is rewritten and no URL is fetched by this script.
"""
import argparse
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import struct
import tempfile
import zipfile
import zlib

ROOT = Path(__file__).resolve().parents[2]
OUT = None
CONFIG = {}
URL = "https://example.invalid/mat06#fonte"
ROWS = [["Item", "Quantidade", "Estado"], ["Alfa", "2", "Conferido"],
        ["Beta", "3", "Conferido"], ["Total", "5", "Conferido"]]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def pinned_file(record):
    candidate = Path(record["path"])
    if not candidate.is_absolute() or not candidate.is_file():
        raise ValueError("Configuration requires an existing absolute file path")
    candidate = candidate.resolve()
    if digest(candidate) != record["sha256"]:
        raise ValueError("Configured file checksum mismatch")
    return candidate


def configure(config_path):
    global CONFIG, OUT
    CONFIG = json.loads(Path(config_path).read_text(encoding="utf-8-sig"))
    candidate = Path(CONFIG["output_dir"])
    OUT = candidate.resolve()
    private = (ROOT / ".private").resolve()
    if not candidate.is_absolute() or OUT == private or not OUT.is_relative_to(private):
        raise ValueError("output_dir must be an absolute path inside checkout .private")
    for original in CONFIG.get("originals", []):
        source = pinned_file(original)
        if source.suffix.lower() not in (".docx", ".pptx") or source.is_relative_to(OUT):
            raise ValueError("Original must be DOCX/PPTX outside output_dir")
    for record in CONFIG.get("tools", {}).values():
        pinned_file(record)


def figure():
    """Deterministic RGB PNG, generated from constants without external assets."""
    width, height = 600, 400
    pixels = bytearray()
    colors = [(35, 78, 112), (220, 125, 50), (50, 140, 95)]
    for y in range(height):
        pixels.append(0)
        for x in range(width):
            band = next((i for i in range(3) if 70 + 90*i <= y < 130 + 90*i), None)
            pixels.extend(colors[band] if band is not None and 80 <= x < 520 else (255, 255, 255))
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
    png += chunk(b"IDAT", zlib.compress(bytes(pixels), 0)) + chunk(b"IEND", b"")
    target = OUT / "synthetic-figure.png"
    target.write_bytes(png)
    return target


def save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def build(docx_only=False):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    FIGURE = figure()
    doc = Document()
    doc.core_properties.author = "AraHub synthetic fixture"
    for name in ("Normal", "Title", "Heading 1"):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Title"].font.size = Pt(24)
    # The bundled default template may carry a blue title rule.
    for border in doc.styles["Title"].element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    doc.add_paragraph("Prova de exportação de materiais", "Title")
    doc.add_paragraph("Este documento sintético permite conferir tabela, hiperlink e figura após a exportação. Não contém dados pessoais.")
    doc.add_paragraph("Tabela de conferência", "Heading 1")
    table = doc.add_table(rows=4, cols=3)
    table.autofit = False
    for col, width in zip(table.columns, (2.7, 1.4, 2.4)):
        col.width = Inches(width)
    for r, row in enumerate(ROWS):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "234E70" if r == 0 else ("EEF3F7" if r % 2 == 0 else "FFFFFF"))
            props.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                m = OxmlElement("w:" + side)
                m.set(qn("w:w"), "100")
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            props.append(margins)
            for run in cell.paragraphs[0].runs:
                run.bold = r == 0
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else "000000")
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "6"), ("color", "D9D9D9")):
            edge.set(qn("w:" + key), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    doc.add_paragraph("Os itens Alfa e Beta somam cinco unidades.")
    p = doc.add_paragraph("Link de teste: ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(URL, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "175899")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = "Referência sintética"
    run.extend([props, text])
    link.append(run)
    p._p.append(link)
    doc.add_paragraph("Figura de conferência", "Heading 1")
    picture = doc.add_picture(str(FIGURE), width=Inches(4.3))
    picture._inline.docPr.set("descr", "Três faixas sintéticas nas cores azul, laranja e verde")
    doc.add_paragraph("Figura 1. Faixas coloridas preservadas como imagem incorporada.")
    doc.save(OUT / "fixture.docx")
    if docx_only:
        return

    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.black
    styles["Heading2"].textColor = colors.black
    styles["BodyText"].fontSize = 11
    pdf = SimpleDocTemplate(str(OUT / "fixture.pdf"), title="Prova de exportação de materiais",
                           author="AraHub synthetic fixture", leftMargin=54, rightMargin=54,
                           topMargin=48, bottomMargin=48)
    pdf_table = Table(ROWS, colWidths=[190, 105, 170], hAlign="LEFT")
    pdf_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#234E70")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#D9D9D9")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("TOPPADDING", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    pdf.build([
        Paragraph("Prova de exportação de materiais", styles["Title"]),
        Paragraph("Este documento sintético permite conferir tabela, hiperlink e figura após a exportação. Não contém dados pessoais.", styles["BodyText"]),
        Spacer(1, 16), Paragraph("Tabela de conferência", styles["Heading2"]), pdf_table,
        Paragraph("Os itens Alfa e Beta somam cinco unidades.", styles["BodyText"]),
        Spacer(1, 10), Paragraph(f'Link de teste: <link href="{URL}" color="#175899"><u>Referência sintética</u></link>', styles["BodyText"]),
        Paragraph("Figura de conferência", styles["Heading2"]),
        Image(str(FIGURE), width=310, height=310*2/3, hAlign="LEFT"),
        Paragraph("Figura 1. Faixas coloridas preservadas como imagem incorporada.", styles["BodyText"]),
    ])
    save("fixtures-inputs.json", {"synthetic": True, "figure_sha256": digest(FIGURE),
         "hyperlink": URL, "rows": ROWS, "source_content_printed": False})


def office_inventory(path):
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        assert sum(i.file_size for i in infos) < 64*1024*1024
        assert all(not i.flag_bits & 1 and i.file_size < 16*1024*1024 for i in infos)
        assert len({i.filename for i in infos}) == len(infos)
        assert all(not i.filename.startswith(("/", "\\")) and ".." not in i.filename.replace("\\", "/").split("/") for i in infos)
        assert not any("vbaproject" in i.filename.lower() or i.filename.lower().endswith(".bin") for i in infos)
        assert archive.testzip() is None
        xml = {i.filename: etree.fromstring(archive.read(i), parser) for i in infos if i.filename.endswith((".xml", ".rels"))}
        for element in xml.values():
            assert not element.getroottree().docinfo.doctype
        hyperlinks = [r.attrib["Target"] for name, root in xml.items() if name.endswith(".rels")
                      for r in root if r.get("Type", "").endswith("/hyperlink")]
        body = [root for name, root in xml.items() if name == "word/document.xml" or
                (name.startswith("ppt/slides/slide") and name.endswith(".xml"))]
        tables = sum(len(root.xpath('//*[local-name()="tbl"]')) for root in body)
        pictures = sum(len(root.xpath('//*[local-name()="pic"]')) for root in body)
        return {"path": str(path), "sha256": digest(path), "bytes": path.stat().st_size,
                "crc_valid": True, "xml_valid": True, "macro_parts": 0,
                "tables": tables, "pictures": pictures, "hyperlinks": hyperlinks,
                "embedded_media": len([i for i in infos if "/media/" in i.filename]),
                "body_parts": len(body)}


def render_docx():
    skill = pinned_file(CONFIG["tools"]["docx_renderer"])
    tool_dirs = [str(pinned_file(CONFIG["tools"][name]).parent)
                 for name in ("soffice", "pdftoppm", "pdfinfo")]
    os.environ["PATH"] = os.pathsep.join(tool_dirs + [os.environ.get("PATH", "")])
    spec = importlib.util.spec_from_file_location("packaged_docx_renderer", skill)
    renderer = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(renderer)
    original_convert = renderer.convert_to_pdf
    def safe_convert(doc_path, profile, *args, **kwargs):
        user = Path(profile) / "user"
        user.mkdir(parents=True, exist_ok=True)
        (user / "registrymodifications.xcu").write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="FirstRun" oor:op="fuse"><value>false</value></prop></item>
</oor:items>''', encoding="utf-8")
        return original_convert(doc_path, profile, *args, **kwargs)
    renderer.convert_to_pdf = safe_convert
    # Private temporary profiles; the managed renderer/dependencies stay unchanged.
    temp = OUT / "tmp"
    temp.mkdir(exist_ok=True)
    tempfile.tempdir = str(temp)
    os.environ["TEMP"] = os.environ["TMP"] = str(temp)
    paths = [(OUT / "fixture.docx", "fixture-docx")]
    for i, original in enumerate(CONFIG.get("originals", []), 1):
        source = pinned_file(original)
        if source.suffix.lower() == ".docx":
            paths.append((source, f"original-{i}-docx"))
    results = []
    for source, label in paths:
        before = digest(source)
        office_inventory(source)  # reject macro parts before invoking Office
        pages = renderer.rasterize(str(source), str(OUT / label), 130, False, True)
        assert pages and digest(source) == before
        results.append({"source_sha256": before, "source_unchanged": True,
                        "pages": [str(Path(p).relative_to(ROOT)) for p in pages],
                        "renderer_sha256": digest(skill),
                        "tools": CONFIG["tools"]})
    save("docx-render.json", results)


def verify():
    from pypdf import PdfReader
    result = {"fixtures": [], "connector_originals": [], "visual_review": "pending"}
    for ext in ("docx", "pptx"):
        report = office_inventory(OUT / ("output/fixture.pptx" if ext == "pptx" else "fixture.docx"))
        assert report["tables"] == 1 and report["pictures"] == 1
        assert report["hyperlinks"] == [URL] and report["embedded_media"] == 1
        result["fixtures"].append(report)
    reader = PdfReader(OUT / "fixture.pdf")
    assert len(reader.pages) == 1
    links = [a.get_object()["/A"]["/URI"] for p in reader.pages for a in p.get("/Annots", [])
             if a.get_object().get("/A", {}).get("/S") == "/URI"]
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert links == [URL] and all(cell in text for row in ROWS for cell in row)
    assert sum(len(p.images) for p in reader.pages) == 1
    result["fixtures"].append({"path": str((OUT / "fixture.pdf").relative_to(ROOT)),
                               "sha256": digest(OUT / "fixture.pdf"), "pages": 1,
                               "table_cells_in_text": True, "hyperlinks": links, "images": 1})
    for original in CONFIG.get("originals", []):
        source = pinned_file(original)
        result["connector_originals"].append(office_inventory(source))
    save("structure.json", result)


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("operation", choices=["build", "build-docx", "render-docx", "verify", "figure", "check-config"])
    p.add_argument("--config", default=os.environ.get("MAT06_CONFIG"))
    args = p.parse_args()
    if not args.config:
        p.error("--config or MAT06_CONFIG is required")
    configure(args.config)
    if args.operation != "check-config":
        OUT.mkdir(parents=True, exist_ok=True)
    {"build": build, "build-docx": lambda: build(True),
     "render-docx": render_docx, "verify": verify, "figure": figure,
     "check-config": lambda: None}[args.operation]()
    print(f"MAT06 {args.operation}: completed; private evidence only")
